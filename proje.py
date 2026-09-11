import streamlit as st
from datetime import date
import smtplib
import io
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from email.mime.text import MIMEText
from email.mime.multipart import MIMEMultipart
from email.mime.application import MIMEApplication

# -------------------------------------------------------------
# Sayfa Yapılandırması & Özel CSS (Koyu Mavi / Beyaz Tema)
# -------------------------------------------------------------
st.set_page_config(page_title="Panthera Lojistik Sipariş Portalı", layout="wide", page_icon="🚚")

st.markdown("""
    <style>
    .main { background-color: #f8f9fa; }
    
    .header-banner {
        background: linear-gradient(135deg, #002060 0%, #004080 100%);
        color: white;
        padding: 30px;
        border-radius: 12px;
        text-align: center;
        margin-bottom: 25px;
        box-shadow: 0 4px 15px rgba(0,0,0,0.15);
    }
    .header-banner h1 {
        color: #ffffff !important;
        font-family: 'Segoe UI', Tahoma, Geneva, Verdana, sans-serif;
        font-weight: 700;
        margin: 0;
        letter-spacing: 1.5px;
    }
    .header-banner p { color: #d1e2ff; margin-top: 8px; font-size: 1.1rem; }

    .section-title {
        color: #002060;
        font-weight: 600;
        border-bottom: 2px solid #002060;
        padding-bottom: 6px;
        margin-top: 20px;
        margin-bottom: 15px;
    }

    .stButton>button {
        background-color: #002060 !important;
        color: white !important;
        font-size: 1.1rem !important;
        font-weight: 600 !important;
        padding: 12px 28px !important;
        border-radius: 8px !important;
        border: none !important;
        transition: all 0.3s ease;
        box-shadow: 0 4px 10px rgba(0,32,96,0.2);
    }
    .stButton>button:hover {
        background-color: #004080 !important;
        box-shadow: 0 6px 15px rgba(0,64,128,0.3);
        transform: translateY(-1px);
    }
    </style>
""", unsafe_allow_html=True)


# -------------------------------------------------------------
# EXCEL OLUŞTURMA FONKSİYONU (Orijinal Şablon Üzerine Yazan Yapı)
# -------------------------------------------------------------
def excel_olustur(siparis_detaylari):
    template_filename_xlsx = "Kopya Sipariş Formu.xlsx"
    template_filename_xls = "Kopya Sipariş Formu.xls"
    
    # Mevcut şablon dosyasını bul ya da sıfırdan aynı görseli oluştur
    if os.path.exists(template_filename_xlsx):
        wb = openpyxl.load_workbook(template_filename_xlsx)
        ws = wb.active
    else:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Boş Form"

        # Görsel Stiller
        header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
        bold_font = Font(name="Calibri", size=10, bold=True)
        blue_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        thin_border = Border(
            left=Side(style='thin', color='000000'),
            right=Side(style='thin', color='000000'),
            top=Side(style='thin', color='000000'),
            bottom=Side(style='thin', color='000000')
        )
        center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
        # Başlık Etiketleri (C2:C4)
        ws.cell(row=2, column=3, value="Sipariş Veren Firma Adı").font = bold_font
        ws.cell(row=3, column=3, value="Yetkili Adı ve Soyadı").font = bold_font
        ws.cell(row=4, column=3, value="Sipariş Tarihi").font = bold_font

        # Kırmızı Uyarı Notu (I2)
        ws.cell(row=2, column=9, value="Palet sayısı ile birlikte kg miktarını da yazmayı unutmayınız.").font = Font(italic=True, bold=True, color="FF0000", size=9)

        # Tablo Başlıkları (Row 5)
        headers = [
            "İRSALİYE NUMARASI",
            "TESLİMAT NOKTASI\n(Bayi ya da zincir mağaza adı )",
            "TESLİMAT NOKTASI YETKİLİ VE İLETİŞİM BİLGİSİ",
            "GİDİLECEK ŞEHİR",
            "AÇIKLAMA",
            "TESLİM TARİHİ",
            "KOLİ SAYISI",
            "KURU PALET SAYISI",
            "SOĞUK PALET SAYISI (+4C)",
            "DONUK PALET SAYISI (-18C)",
            "KURU KG",
            "SOĞUK KG",
            "DONUK KG"
        ]
        for col_idx, header in enumerate(headers, start=3):
            c = ws.cell(row=5, column=col_idx, value=header)
            c.font = header_font
            c.fill = blue_fill
            c.alignment = center_align
            c.border = thin_border

    # --- Şablon Üzerine Kullanıcı Verilerini Doldurma ---
    regular_font = Font(name="Calibri", size=10)
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    center_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Üst bilgiler (D2, D3, D4 hücreleri)
    ws.cell(row=2, column=4, value=siparis_detaylari['Firma']).font = regular_font
    ws.cell(row=3, column=4, value=siparis_detaylari['Yetkili']).font = regular_font
    ws.cell(row=4, column=4, value=str(date.today())).font = regular_font

    # Sipariş Satırı Verileri (6. Satır, C6 - O6 arası)
    teslimat_noktasi = f"{siparis_detaylari['Depo Kodu']} - {siparis_detaylari['Adres']}"
    row_values = [
        siparis_detaylari['İrsaliye'],
        teslimat_noktasi,
        siparis_detaylari['Yetkili Tel'],
        f"{siparis_detaylari['İlce']} / {siparis_detaylari['İl']}",
        siparis_detaylari['Açıklama'],
        siparis_detaylari['Teslim Tarihi'],
        siparis_detaylari['Koli'],
        siparis_detaylari['Kuru Palet'],
        siparis_detaylari['Soğuk Palet'],
        siparis_detaylari['Donuk Palet'],
        siparis_detaylari['Kuru KG'],
        siparis_detaylari['Soğuk KG'],
        siparis_detaylari['Donuk KG']
    ]

    for col_idx, val in enumerate(row_values, start=3):
        cell = ws.cell(row=6, column=col_idx, value=val)
        cell.font = regular_font
        cell.border = thin_border
        cell.alignment = center_align

    excel_buffer = io.BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer.getvalue()


# -------------------------------------------------------------
# E-POSTA GÖNDERİM FONKSİYONU
# -------------------------------------------------------------
def siparis_mailleri_gonder(siparis_detaylari, musteri_epostaları_raw):
    SMTP_SUNUCU = "smtp.gmail.com"
    SMTP_PORT = 587

    SISTEM_GONDERICI_HESAPLAR = [
        {"eposta": "yildizsususu@gmail.com", "sifre": "frbfqtcneiyoqpre"},
    ]

    HEDEF_OPERASYON_ALICILARI = [
        "sudeozkoc@pantheralojistik.com.tr",
    ]

    musteri_eposta_listesi = [
        e.strip() for e in musteri_epostaları_raw.replace(";", ",").split(",") if e.strip()
    ]

    excel_data = excel_olustur(siparis_detaylari)
    excel_dosya_adi = f"{siparis_detaylari['Firma']}_Siparisi_{siparis_detaylari['İrsaliye']}.xlsx"

    konu_yonetici = f"🚨 {siparis_detaylari['Firma']} Siparişi - İrsaliye No: {siparis_detaylari['İrsaliye']}"
    konu_musteri = f"✅ {siparis_detaylari['Firma']} Siparişi Alındı - İrsaliye No: {siparis_detaylari['İrsaliye']}"

    body_yonetici = f"""
    <div style="font-family: Arial, sans-serif; color: #333;">
        <h2 style="color: #002060;">Panthera Lojistik - {siparis_detaylari['Firma']} Siparişi</h2>
        <p><b>{siparis_detaylari['Firma']}</b> firmasından yeni bir sipariş kaydı oluşturuldu. Orijinal form formatındaki Excel dosyası ektedir.</p>
        <hr style="border: 1px solid #002060;">
        <ul>
            <li><b>Sipariş Veren Firma:</b> {siparis_detaylari['Firma']}</li>
            <li><b>Yetkili Adı Soyadı:</b> {siparis_detaylari['Yetkili']}</li>
            <li><b>Müşteri İletişim E-Posta(ları):</b> {', '.join(musteri_eposta_listesi)}</li>
            <li><b>İrsaliye Numarası:</b> {siparis_detaylari['İrsaliye']}</li>
            <li><b>Planlanan Teslim Tarihi:</b> {siparis_detaylari['Teslim Tarihi']}</li>
        </ul>
        <br>
        <p><i>Detaylı palet, adres ve koli bilgileri için ekteki Excel dosyasını inceleyebilirsiniz.</i></p>
    </div>
    """

    body_musteri = f"""
    <div style="font-family: Arial, sans-serif; color: #333;">
        <h2 style="color: #002060;">Panthera Lojistik Sipariş Portalı</h2>
        <p>Sayın <b>{siparis_detaylari['Yetkili']}</b>,</p>
        <p><b>{siparis_detaylari['Firma']}</b> adına oluşturmuş olduğunuz sipariş talebiniz operasyon ekibimize başarıyla ulaşmıştır. Siparişinizin Excel formatındaki detay örneği ektedir.</p>
        <hr style="border: 1px solid #002060;">
        <p><b>Sipariş Özetiniz:</b></p>
        <ul>
            <li><b>İrsaliye No:</b> {siparis_detaylari['İrsaliye']}</li>
            <li><b>Teslimat Adresi:</b> {siparis_detaylari['İlce']} / {siparis_detaylari['İl']}</li>
            <li><b>Teslim Tarihi:</b> {siparis_detaylari['Teslim Tarihi']}</li>
            <li><b>Toplam Palet:</b> {siparis_detaylari['Kuru Palet'] + siparis_detaylari['Soğuk Palet'] + siparis_detaylari['Donuk Palet']} Palet</li>
        </ul>
        <p>Bizi tercih ettiğiniz için teşekkür ederiz.</p>
        <br>
        <p style="font-size: 0.9rem; color: #666;"><i>Panthera Lojistik Operasyon Ekibi</i></p>
    </div>
    """

    last_error = ""

    for gonderici in SISTEM_GONDERICI_HESAPLAR:
        try:
            server = smtplib.SMTP(SMTP_SUNUCU, SMTP_PORT)
            server.starttls()
            server.login(gonderici["eposta"], gonderici["sifre"])

            # 1. Mail: Operasyon Ekibine
            msg1 = MIMEMultipart()
            msg1['From'] = f"Panthera Sipariş Portalı <{gonderici['eposta']}>"
            msg1['To'] = ", ".join(HEDEF_OPERASYON_ALICILARI)
            msg1['Reply-To'] = ", ".join(musteri_eposta_listesi)
            msg1['Subject'] = konu_yonetici
            msg1.attach(MIMEText(body_yonetici, 'html'))

            attachment1 = MIMEApplication(excel_data, Name=excel_dosya_adi)
            attachment1['Content-Disposition'] = f'attachment; filename="{excel_dosya_adi}"'
            msg1.attach(attachment1)

            server.send_message(msg1)

            # 2. Mail: Müşteriye
            msg2 = MIMEMultipart()
            msg2['From'] = f"Panthera Lojistik <{gonderici['eposta']}>"
            msg2['To'] = ", ".join(musteri_eposta_listesi)
            msg2['Subject'] = konu_musteri
            msg2.attach(MIMEText(body_musteri, 'html'))

            attachment2 = MIMEApplication(excel_data, Name=excel_dosya_adi)
            attachment2['Content-Disposition'] = f'attachment; filename="{excel_dosya_adi}"'
            msg2.attach(attachment2)

            server.send_message(msg2)

            server.quit()
            return True, "Başarılı"
        except Exception as e:
            last_error = str(e)
            continue

    return False, last_error


# -------------------------------------------------------------
# HAZIR DEPO VERİTABANI
# -------------------------------------------------------------
HAZIR_DEPOLAR = [
    {
        "kod": "001 - BİM / Konya Depo (Karatay)",
        "il": "Konya",
        "ilce": "Karatay",
        "adres": "Fevzi Çakmak Mah. 10542. Sk. No:12",
        "posta_kodu": "42050"
    },
    {
        "kod": "002 - BİM / Konya Depo (Selçuklu)",
        "il": "Konya",
        "ilce": "Selçuklu",
        "adres": "Organize Sanayi Bölgesi 5. Sok No:8",
        "posta_kodu": "42250"
    },
    {
        "kod": "003 - A101 / Ankara Lojistik Merkezi",
        "il": "Ankara",
        "ilce": "Kazan",
        "adres": "Saray Mah. 100. Yıl Bulvarı No:45",
        "posta_kodu": "06980"
    }
]

# -------------------------------------------------------------
# ARAYÜZ (FRONTEND)
# -------------------------------------------------------------

st.markdown("""
    <div class="header-banner">
        <h1>PANTHERA LOJİSTİK SİPARİŞ PORTALI</h1>
        <p>Lütfen sipariş ve teslimat bilgilerinizi eksiksiz doldurunuz.</p>
    </div>
""", unsafe_allow_html=True)

# 1. BÖLÜM: MÜŞTERİ KİMLİK BİLGİLERİ
st.markdown('<h3 class="section-title">1. Müşteri ve İletişim Bilgileri</h3>', unsafe_allow_html=True)

col_f1, col_f2, col_f3 = st.columns(3)
with col_f1:
    siparis_veren = st.text_input("Sipariş Veren Firma Adı *")
with col_f2:
    yetkili_ad = st.text_input("Sipariş Veren Yetkili Adı Soyadı *")
with col_f3:
    musteri_eposta = st.text_input("E-Posta Adresiniz (Birden fazla ise virgül ile ayırınız) *", help="Örn: ahmet@firma.com, mehmet@firma.com")

# 2. BÖLÜM: TESLİMAT NOKTASI VE ADRES BİLGİLERİ
st.markdown('<h3 class="section-title">2. Teslimat Noktası ve Adres Seçimi</h3>', unsafe_allow_html=True)

depo_secenekleri = ["Seçiniz..."] + [d["kod"] for d in HAZIR_DEPOLAR] + ["➕ Listede Yok / Yeni Adres Gireceğim"]

secilen_depo_kodu = st.selectbox(
    "Teslimat Yapılacak Depo veya Adres *",
    options=depo_secenekleri,
    index=0
)

secilen_hazir_depo = next((d for d in HAZIR_DEPOLAR if d["kod"] == secilen_depo_kodu), None)

if secilen_depo_kodu == "➕ Listede Yok / Yeni Adres Gireceğim":
    st.info("💡 Aradığınız depo listede yoksa aşağıdaki alanlara teslimat yapılacak açık adresi eksiksiz giriniz.")
    
    col_a1, col_a2, col_a3, col_a4 = st.columns(4)
    with col_a1:
        il = st.text_input("İl (Şehir) *")
    with col_a2:
        ilce = st.text_input("İlçe *")
    with col_a3:
        posta_kodu = st.text_input("Posta Kodu *")
    with col_a4:
        yetkili_tel = st.text_input("Teslimat Noktası Yetkili/Tel")
    
    acik_adres = st.text_area("Açık Adres (Cadde, Mahalle, Sokak, No) *")

else:
    col_a1, col_a2, col_a3, col_a4 = st.columns(4)
    with col_a1:
        il = st.text_input("İl (Şehir) *", value=secilen_hazir_depo["il"] if secilen_hazir_depo else "", disabled=True)
    with col_a2:
        ilce = st.text_input("İlçe *", value=secilen_hazir_depo["ilce"] if secilen_hazir_depo else "", disabled=True)
    with col_a3:
        posta_kodu = st.text_input("Posta Kodu *", value=secilen_hazir_depo["posta_kodu"] if secilen_hazir_depo else "", disabled=True)
    with col_a4:
        yetkili_tel = st.text_input("Teslimat Noktası Yetkili/Tel")
    
    acik_adres = st.text_area("Açık Adres *", value=secilen_hazir_depo["adres"] if secilen_hazir_depo else "", disabled=True)

# 3. BÖLÜM: YÜK VE İRSALİYE DETAYLARI
st.markdown('<h3 class="section-title">3. Yük ve İrsaliye Detayları</h3>', unsafe_allow_html=True)

col_d1, col_d2, col_d3 = st.columns(3)

with col_d1:
    irsaliye_no = st.text_input("İrsaliye Numarası (İlk 3 Harf + 13 Rakam = 16 Karakter) *", max_chars=16, help="Örn: ABC2026000001234")
    teslim_tarihi = st.date_input("Planlanan Teslim Tarihi *", value=date.today())
    aciklama = st.text_input("Açıklama / Özel Notlar")

with col_d2:
    koli_sayisi = st.number_input("Koli Sayısı", min_value=0, step=1, value=0)
    kuru_palet = st.number_input("Kuru Palet Sayısı", min_value=0, step=1, value=0)
    soguk_palet = st.number_input("Soğuk Palet (+4°C) Sayısı", min_value=0, step=1, value=0)
    donuk_palet = st.number_input("Donuk Palet (-18°C) Sayısı", min_value=0, step=1, value=0)

with col_d3:
    kuru_kg = st.number_input("Kuru Yük KG *", min_value=0.0, step=10.0, value=0.0)
    soguk_kg = st.number_input("Soğuk Yük (+4°C) KG *", min_value=0.0, step=10.0, value=0.0)
    donuk_kg = st.number_input("Donuk Yük (-18°C) KG *", min_value=0.0, step=10.0, value=0.0)

st.markdown("<br>", unsafe_allow_html=True)

# -------------------------------------------------------------
# GÖNDER BUTONU VE KONTROLLER
# -------------------------------------------------------------
if st.button("🚀 Siparişi Onayla ve Gönder", use_container_width=True):
    irsaliye_clean = irsaliye_no.strip()
    
    # 1. Zorunlu Alan Kontrolü
    if not siparis_veren or not yetkili_ad or not musteri_eposta or not irsaliye_clean or secilen_depo_kodu == "Seçiniz...":
        st.error("❌ Lütfen kırmızı yıldızlı (*) zorunlu alanları (Firma, Yetkili, E-Posta, İrsaliye No, Teslimat Noktası) doldurunuz!")
    
    # 2. İrsaliye Numarası Format Kontrolü (Tam 16 Karakter ve İlk 3'ü Harf)
    elif len(irsaliye_clean) != 16 or not irsaliye_clean[:3].isalpha():
        st.error("❌ İrsaliye Numarası geçersiz! İrsaliye No **tam 16 karakter** olmalı ve **ilk 3 karakteri harf** içermelidir (Örn: ABC2026000001234).")
    
    # 3. Yeni Adres Kontrolü
    elif secilen_depo_kodu == "➕ Listede Yok / Yeni Adres Gireceğim" and (not il or not ilce or not acik_adres or not posta_kodu):
        st.error("❌ Yeni adres seçeneğini seçtiniz. Lütfen Şehir, İlçe, Açık Adres ve Posta Kodu alanlarını doldurunuz!")
    
    # 4. Yük Miktarı Genel Kontrolü
    elif (kuru_palet + soguk_palet + donuk_palet) == 0 and koli_sayisi == 0:
        st.warning("⚠️ Lütfen en az bir adet Palet veya Koli miktarı giriniz!")

    # 5. ZORUNLU KG KONTROLÜ (Palet sayısı girilmişse KG miktarı 0 olamaz)
    elif kuru_palet > 0 and kuru_kg <= 0:
        st.error("❌ Kuru Palet girdiniz. Lütfen **Kuru Yük KG** miktarını da yazınız!")
    elif soguk_palet > 0 and soguk_kg <= 0:
        st.error("❌ Soğuk Palet (+4°C) girdiniz. Lütfen **Soğuk Yük KG** miktarını da yazınız!")
    elif donuk_palet > 0 and donuk_kg <= 0:
        st.error("❌ Donuk Palet (-18°C) girdiniz. Lütfen **Donuk Yük KG** miktarını da yazınız!")
    
    else:
        siparis_verileri = {
            "Firma": siparis_veren,
            "Yetkili": yetkili_ad,
            "İrsaliye": irsaliye_clean,
            "Depo Kodu": secilen_depo_kodu,
            "İl": il,
            "İlce": ilce,
            "Posta Kodu": posta_kodu,
            "Adres": acik_adres,
            "Yetkili Tel": yetkili_tel if yetkili_tel else "Belirtilmedi",
            "Teslim Tarihi": str(teslim_tarihi),
            "Koli": koli_sayisi,
            "Kuru Palet": kuru_palet,
            "Soğuk Palet": soguk_palet,
            "Donuk Palet": donuk_palet,
            "Kuru KG": kuru_kg,
            "Soğuk KG": soguk_kg,
            "Donuk KG": donuk_kg,
            "Açıklama": aciklama if aciklama else ""
        }
        
        with st.spinner("Sipariş işleniyor ve Excel dosyası e-postaya ekleniyor..."):
            basari, hata_mesaji = siparis_mailleri_gonder(siparis_verileri, musteri_eposta)
        
        if basari:
            st.success("🎉 Siparişiniz Excel dosyası olarak operasyon ekibine ve e-posta adreslerinize başarıyla gönderildi!")
            st.balloons()
        else:
            st.error(f"❌ E-posta gönderilemedi! Hata Detayı: {hata_mesaji}")
            st.info("💡 **İpucu:** Lütfen `SISTEM_GONDERICI_HESAPLAR` alanındaki Gmail şifrenizin normal şifre değil, Google 'Uygulama Şifresi' olduğunu doğrulayın.")
